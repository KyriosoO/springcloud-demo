from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document
import pymupdf
import pytest

from knowledge_corpus_tools.errors import ContractError
from knowledge_corpus_tools.models import BlockKind
from knowledge_corpus_tools.normalization_v2 import TrustStatus, parsed_source_v2
from knowledge_corpus_tools.parsing import parse_asset, parse_asset_v2


def _docx_with_interleaved_table(path: Path) -> str:
    document = Document()
    document.add_paragraph("第一条 一般规定")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "条件"
    document.add_paragraph("第二条 例外规定")
    document.save(str(path))
    return sha256(path.read_bytes()).hexdigest()


def test_docx_v2_preserves_interleaved_body_order_without_changing_v1(tmp_path: Path) -> None:
    path = tmp_path / "source.docx"
    digest = _docx_with_interleaved_table(path)

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    ordered = parse_asset_v2(path, asset_id="source", asset_sha256=digest)

    assert legacy.parser_version == "1.2.0"
    assert [block.kind for block in legacy.blocks] == [BlockKind.CLAUSE, BlockKind.CLAUSE, BlockKind.TABLE]
    assert ordered.parser_version == "1.2.0+ordered-v2"
    assert [block.kind for block in ordered.blocks] == [BlockKind.CLAUSE, BlockKind.TABLE, BlockKind.CLAUSE]
    assert [block.ordinal for block in ordered.blocks] == [1, 2, 3]
    assert ordered.blocks[1].table_id == "table-1"
    assert ordered.blocks[1].text == "项目 | 条件"
    assert ordered.blocks[2].clause_id == "第二条"


def test_docx_v2_order_proof_does_not_claim_flattened_table_structure(tmp_path: Path) -> None:
    path = tmp_path / "source.docx"
    digest = _docx_with_interleaved_table(path)
    ordered = parse_asset_v2(path, asset_id="source", asset_sha256=digest)

    source = parsed_source_v2(ordered, order_proof_ref=f"source-docx-sha256:{digest}")
    assert source.order_status is TrustStatus.VERIFIED
    assert source.structure_status is TrustStatus.REVIEW_REQUIRED
    assert source.text == "第一条 一般规定\n项目 | 条件\n第二条 例外规定"


def test_parse_asset_v2_preserves_other_parser_behavior(tmp_path: Path) -> None:
    path = tmp_path / "source.html"
    path.write_text("<html><body><p>税务政策</p></body></html>", encoding="utf-8")
    digest = sha256(path.read_bytes()).hexdigest()

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert legacy.parser_version == "4.15.0"
    assert revised.parser_version == "4.15.0+sections-v2"
    assert revised.blocks == legacy.blocks


@pytest.mark.parametrize("suffix", [".html", ".docx"])
def test_parse_asset_v2_rejects_mismatched_asset_bytes_before_parsing(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"source{suffix}"
    if suffix == ".docx":
        _docx_with_interleaved_table(path)
    else:
        path.write_text("<p>正文</p>", encoding="utf-8")

    with pytest.raises(ContractError, match="asset bytes SHA-256 mismatch"):
        parse_asset_v2(path, asset_id="source", asset_sha256="0" * 64)
    assert parse_asset(path, asset_id="source", asset_sha256="0" * 64).asset_sha256 == "0" * 64


def test_docx_v2_recognizes_sections_and_clause_overrides_heading_style(tmp_path: Path) -> None:
    path = tmp_path / "sections.docx"
    document = Document()
    document.add_paragraph("第一节 适用范围")
    document.add_paragraph("第一条 标题样式", style="Heading 1")
    document.add_paragraph("第一节中规定的条件")
    document.save(str(path))
    digest = sha256(path.read_bytes()).hexdigest()

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert [block.kind for block in legacy.blocks] == [BlockKind.PARAGRAPH, BlockKind.HEADING, BlockKind.PARAGRAPH]
    assert [block.kind for block in revised.blocks] == [BlockKind.HEADING, BlockKind.CLAUSE, BlockKind.PARAGRAPH]
    assert revised.blocks[1].clause_id == "第一条"


def test_legacy_doc_v2_recognizes_sections_without_changing_v1(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class Result:
        text = "第一节 一般性规定\n第一条 具体规定\n第一节中规定的条件"
        version = "0.2.1"

    import legacy_doc
    import knowledge_corpus_tools.parsing as parsing

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(
        parsing, "_legacy_doc_raw_text",
        lambda _: "第一节 一般性规定\r第一条 具体规定\r第一节中规定的条件",
    )
    path = tmp_path / "sections.doc"
    path.write_bytes(b"synthetic legacy doc")
    digest = sha256(path.read_bytes()).hexdigest()

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert legacy.parser_version == "0.2.1+structure-v1"
    assert revised.parser_version == "0.2.1+native-paragraphs-v2"
    assert [block.kind for block in legacy.blocks] == [BlockKind.PARAGRAPH, BlockKind.CLAUSE, BlockKind.PARAGRAPH]
    assert [block.kind for block in revised.blocks] == [BlockKind.HEADING, BlockKind.CLAUSE, BlockKind.PARAGRAPH]
    assert revised.blocks[1].section_path == ("第一节 一般性规定",)
    assert revised.blocks[2].section_path == ("第一节 一般性规定",)


def test_legacy_doc_v2_keeps_manual_line_breaks_within_one_native_paragraph(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import legacy_doc
    import knowledge_corpus_tools.parsing as parsing
    from legacy_doc.normalize import normalize_word_text

    native = ("第一节 范围\r公办机构符合条件。\x0b民办机构符合条件。"
              "\x0b超过标准的部分不适用。\r金融机构是指：\x0b（1）银行。\x0b（2）信用合作社。")

    class Result:
        text = normalize_word_text(native)
        version = "0.2.1"

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(parsing, "_legacy_doc_raw_text", lambda _: native)
    path = tmp_path / "manual-breaks.doc"
    path.write_bytes(b"synthetic legacy doc")
    digest = sha256(path.read_bytes()).hexdigest()

    old = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert len(old.blocks) == 7
    assert len(revised.blocks) == 3
    assert revised.blocks[1].kind is BlockKind.PARAGRAPH
    assert revised.blocks[1].text == (
        "公办机构符合条件。\n民办机构符合条件。\n超过标准的部分不适用。"
    )
    assert revised.blocks[2].text == "金融机构是指：\n（1）银行。\n（2）信用合作社。"
    source = parsed_source_v2(revised)
    assert source.order_status is TrustStatus.UNVERIFIED
    assert source.structure_status is TrustStatus.UNVERIFIED
    assert source.blocks[1].end - source.blocks[1].start == len(revised.blocks[1].text)


def test_legacy_doc_v2_rejects_public_text_or_version_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import legacy_doc
    import knowledge_corpus_tools.parsing as parsing

    class Result:
        text = "正文有遗漏"
        version = "0.2.1"

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(parsing, "_legacy_doc_raw_text", lambda _: "完整正文\r")
    path = tmp_path / "drift.doc"
    path.write_bytes(b"synthetic legacy doc")
    digest = sha256(path.read_bytes()).hexdigest()
    with pytest.raises(ContractError, match="differ from public extraction"):
        parse_asset_v2(path, asset_id="source", asset_sha256=digest)

    monkeypatch.setattr(parsing, "version", lambda _: "0.2.0")
    with pytest.raises(ContractError, match="unsupported legacy DOC parser version"):
        parse_asset_v2(path, asset_id="source", asset_sha256=digest)


@pytest.mark.parametrize(
    ("native", "expected_error"),
    [
        ("项目\x07金额", "table structure requires review"),
        ("第一节 范围\x0b正文条件。", "mixed native paragraph structure requires review"),
    ],
)
def test_legacy_doc_v2_rejects_unprovable_native_structure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, native: str, expected_error: str,
) -> None:
    import legacy_doc
    import knowledge_corpus_tools.parsing as parsing
    from legacy_doc.normalize import normalize_word_text

    class Result:
        text = normalize_word_text(native)
        version = "0.2.1"

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(parsing, "_legacy_doc_raw_text", lambda _: native)
    path = tmp_path / "unprovable.doc"
    path.write_bytes(b"synthetic legacy doc")
    with pytest.raises(ContractError, match=expected_error):
        parse_asset_v2(path, asset_id="source", asset_sha256=sha256(path.read_bytes()).hexdigest())


def test_legacy_doc_v2_preserves_hyperlink_display_not_word_field_code(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import legacy_doc
    import knowledge_corpus_tools.parsing as parsing
    from legacy_doc.normalize import normalize_word_text

    native = ('政策说明\x13 HYPERLINK "https://example.invalid/policy" \\t "_blank"'
              '\x14适用范围\x15仍需核实。')

    class Result:
        text = normalize_word_text(native)
        version = "0.2.1"

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(parsing, "_legacy_doc_raw_text", lambda _: native)
    path = tmp_path / "hyperlink.doc"
    path.write_bytes(b"synthetic legacy doc")

    revised = parse_asset_v2(path, asset_id="source", asset_sha256=sha256(path.read_bytes()).hexdigest())
    assert [block.text for block in revised.blocks] == ["政策说明 适用范围仍需核实。"]


def _synthetic_doc_pieces(
    *, cps: tuple[int, int, int] = (0, 2, 4),
    fcs: tuple[int, int] = (0x40, 0x48),
    second_text: bytes = "丙丁".encode("utf-16le"),
) -> tuple[bytes, bytes]:
    word = bytearray(0x80)
    word[0x18:0x1C] = (0x40).to_bytes(4, "little")
    word[0x1C:0x20] = (0x50).to_bytes(4, "little")
    word[0x40:0x44] = "甲乙".encode("utf-16le")
    word[0x48:0x48 + len(second_text)] = second_text
    clx = b"".join(cp.to_bytes(4, "little") for cp in cps)
    clx += b"".join(b"\0\0" + fc.to_bytes(4, "little") + b"\0\0" for fc in fcs)
    return bytes(word), clx


def test_legacy_doc_v2_rejects_in_memory_second_piece_outside_fc_mac(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import legacy_doc
    import legacy_doc.ole
    import legacy_doc.word

    word, clx = _synthetic_doc_pieces()

    class Reader:
        def __init__(self, _: bytes, *, options: Any) -> None:
            pass

        def read_stream(self, name: str) -> bytes:
            return word if name == "WordDocument" else b"table"

    class Result:
        text = "甲乙丙丁"
        version = "0.2.1"

    monkeypatch.setattr(legacy_doc, "extract_text", lambda _: Result())
    monkeypatch.setattr(legacy_doc.ole, "OleReader", Reader)
    monkeypatch.setattr(legacy_doc.word, "_find_clx", lambda _: clx)
    path = tmp_path / "pieces.doc"
    path.write_bytes(b"synthetic legacy doc")
    digest = sha256(path.read_bytes()).hexdigest()
    assert [block.text for block in parse_asset_v2(path, asset_id="source", asset_sha256=digest).blocks] == ["甲乙丙丁"]

    tampered = bytearray(clx)
    # The second PCD still points inside WordDocument, but beyond fcMac.
    tampered[22:26] = (0x60).to_bytes(4, "little")
    monkeypatch.setattr(legacy_doc.word, "_find_clx", lambda _: bytes(tampered))
    with pytest.raises(ContractError, match="outside the complete source range"):
        parse_asset_v2(path, asset_id="source", asset_sha256=digest)


@pytest.mark.parametrize(
    ("cps", "fcs", "second_text", "expected_error"),
    [
        ((1, 3, 5), (0x40, 0x48), "丙丁".encode("utf-16le"), "positions have gaps"),
        ((0, 2, 2), (0x40, 0x48), "丙丁".encode("utf-16le"), "empty pieces"),
        ((0, 2, 4), (0x40, 0x48), b"\x00\xd8\x00\x00", "undecodable text"),
        ((0, 2, 4), (0x40, 0x48), "\ufffd丁".encode("utf-16le"), "replacement characters"),
    ],
)
def test_legacy_doc_v2_piece_preflight_rejects_silent_loss(
    cps: tuple[int, int, int], fcs: tuple[int, int], second_text: bytes, expected_error: str,
) -> None:
    import knowledge_corpus_tools.parsing as parsing

    word, clx = _synthetic_doc_pieces(cps=cps, fcs=fcs, second_text=second_text)
    with pytest.raises(ContractError, match=expected_error):
        parsing._validate_legacy_doc_pieces(word, clx, fc_min=0x40, fc_mac=0x50)


def test_legacy_doc_v2_piece_preflight_keeps_valid_compressed_cp1252_and_rejects_replacement() -> None:
    import knowledge_corpus_tools.parsing as parsing

    word, clx = _synthetic_doc_pieces()
    compressed_word = bytearray(word)
    compressed_clx = bytearray(clx)
    compressed_word[0x40:0x42] = b"AB"
    compressed_clx[14:18] = (0x40000000 | (0x40 * 2)).to_bytes(4, "little")
    parsing._validate_legacy_doc_pieces(bytes(compressed_word), bytes(compressed_clx), fc_min=0x40, fc_mac=0x50)

    compressed_word[0x41] = 0x81  # Undefined in strict cp1252; the pinned reader would replace it.
    with pytest.raises(ContractError, match="undecodable text"):
        parsing._validate_legacy_doc_pieces(bytes(compressed_word), bytes(compressed_clx), fc_min=0x40, fc_mac=0x50)


def test_html_v2_recognizes_only_bounded_section_heading(tmp_path: Path) -> None:
    path = tmp_path / "sections.html"
    path.write_text(
        "<html><body><p>第一节 适用范围</p><p>第一条 具体规定</p>"
        "<p>第一节中规定的条件</p></body></html>", encoding="utf-8",
    )
    digest = sha256(path.read_bytes()).hexdigest()

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert [block.kind for block in legacy.blocks] == [BlockKind.PARAGRAPH, BlockKind.CLAUSE, BlockKind.PARAGRAPH]
    assert [block.kind for block in revised.blocks] == [BlockKind.HEADING, BlockKind.CLAUSE, BlockKind.PARAGRAPH]


def test_html_v2_clause_text_overrides_heading_tag(tmp_path: Path) -> None:
    path = tmp_path / "article.html"
    path.write_text("<html><body><h2>第一条 适用范围</h2></body></html>", encoding="utf-8")
    digest = sha256(path.read_bytes()).hexdigest()
    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert legacy.blocks[0].kind is BlockKind.HEADING
    assert revised.blocks[0].kind is BlockKind.CLAUSE
    assert revised.blocks[0].clause_id == "第一条"


def test_real_pdf_v2_recognizes_section_heading(tmp_path: Path) -> None:
    path = tmp_path / "sections.pdf"
    pdf_api: Any = pymupdf
    document: Any = pdf_api.open()
    page = document.new_page()
    page.insert_font(fontname="CJK", fontbuffer=pdf_api.Font("china-s").buffer)
    page.insert_text((72, 72), "第一节 适用范围", fontname="CJK", fontsize=12)
    page.insert_text((72, 100), "第一条 具体规定", fontname="CJK", fontsize=12)
    path.write_bytes(document.tobytes())
    document.close()
    digest = sha256(path.read_bytes()).hexdigest()

    legacy = parse_asset(path, asset_id="source", asset_sha256=digest)
    revised = parse_asset_v2(path, asset_id="source", asset_sha256=digest)
    assert legacy.parser_version == "6.16.2+structure-v1"
    assert revised.parser_version == "6.16.2+sections-v2"
    assert [block.kind for block in legacy.blocks] == [BlockKind.PARAGRAPH, BlockKind.CLAUSE]
    assert [block.kind for block in revised.blocks] == [BlockKind.HEADING, BlockKind.CLAUSE]
    assert revised.blocks[0].text == "第一节 适用范围"
